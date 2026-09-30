#!/usr/bin/env python3
"""검토본 docx에 붙여 넣은 주제별 설명을 PsyKNLE/TopicGuides.swift로 옮긴다.

검토본(PsyKNLE_답안_검토본_v2.docx)에서 각 문항의 【국시 빈출】 설명 바로 뒤에
추가된 문단 묶음을 찾아, 그 문항의 topic에 대한 설명으로 저장한다.
문단 서식은 다음처럼 옮긴다.
  - 전체가 굵은 문단(목록 아님) → "# 소제목"
  - 목록 문단 → 수준(ilvl)에 따라 "- ", "  - ", "    - "
  - 부분 굵게(run 서식 또는 붙여 넣으며 남은 **…**) → 마크다운 **…**

사용법:
  python3 Scripts/import_topic_guides.py 원본검토본.docx 수정검토본.docx
"""

from __future__ import annotations

import difflib
import re
import sys
from pathlib import Path

import docx

REPO_ROOT = Path(__file__).resolve().parent.parent
APP_FILE = REPO_ROOT / "PsyKNLE" / "PsyKNLEApp.swift"
OUT_FILE = REPO_ROOT / "PsyKNLE" / "TopicGuides.swift"

# 설명 하나가 여러 주제를 함께 다루는 경우: 설명이 붙은 주제 → 함께 보여 줄 주제
SHARED_TOPICS = {
    "성격장애의 유형": ["성격장애 환자의 중재"],
    "물질 관련 장애": ["알코올 관련 장애"],
}

MD_SPECIAL = re.compile(r"([\\`*_\[\]])")


def question_topics() -> dict[str, str]:
    src = APP_FILE.read_text(encoding="utf-8")
    return {m.group(1): m.group(2) for m in re.finditer(
        r'id: "(\d{4}-\d+)".*?topic: "([^"]*)"', src, re.S)}


def paragraph_lines(p) -> list[tuple[str, bool]]:
    """문단을 줄바꿈(<w:br/>) 기준으로 나누고, 줄마다 마크다운 문자열과
    '줄 전체가 굵은지'를 돌려준다."""
    chars: list[tuple[str, bool]] = []
    for r in p.runs:
        bold = bool(r.bold)
        for ch in r.text:
            chars.append((ch, bold))
    lines: list[list[tuple[str, bool]]] = [[]]
    for ch, b in chars:
        if ch == "\n":
            lines.append([])
        else:
            lines[-1].append((ch, b))
    out = []
    for line in lines:
        text = "".join(c for c, _ in line)
        mask = [b for _, b in line]
        # 붙여 넣으며 글자로 남은 **…** 는 굵게 표시로 바꾼다.
        t, m, i, on = [], [], 0, False
        while i < len(text):
            if text.startswith("**", i):
                on = not on
                i += 2
                continue
            t.append(text[i])
            m.append(mask[i] or on)
            i += 1
        text = "".join(t).rstrip()
        m = m[: len(text)]
        visible = [b for c, b in zip(text, m) if c.strip()]
        all_bold = bool(visible) and all(visible)
        # 공백은 앞뒤 글자가 모두 굵을 때만 굵게 친다. 그래야 "**2026-71 **"처럼
        # 닫는 ** 앞에 공백이 끼어 굵게 그려지지 않는 일이 없다.
        eff = list(m)
        for k, c in enumerate(text):
            if not c.strip():
                prev = next((m[j] for j in range(k - 1, -1, -1) if text[j].strip()), False)
                nxt = next((m[j] for j in range(k + 1, len(text)) if text[j].strip()), False)
                eff[k] = prev and nxt
        md, cur = [], False
        for c, b in zip(text, eff):
            if b != cur:
                md.append("**")
                cur = b
            md.append(MD_SPECIAL.sub(r"\\\1", c))
        if cur:
            md.append("**")
        s = "".join(md)
        out.append((s.strip(), all_bold, "".join(t).strip()))
    return out


def list_level(p) -> int | None:
    pPr = p._p.pPr
    if pPr is None or pPr.numPr is None:
        return None
    ilvl = pPr.numPr.ilvl
    return int(ilvl.val) if ilvl is not None else 0


def convert(paragraphs) -> tuple[str, list[str]]:
    lines: list[str] = []
    title = ""
    for p in paragraphs:
        level = list_level(p)
        for i, (md, all_bold, plain) in enumerate(paragraph_lines(p)):
            if not plain:
                if lines and lines[-1] != "":
                    lines.append("")
                continue
            if not title:
                title = plain
                continue
            if level is not None and i == 0:
                lines.append("  " * level + "- " + md)
            elif level is not None:
                lines.append("  " * (level + 1) + md)
            elif all_bold:
                if lines and lines[-1] != "":
                    lines.append("")
                lines.append("# " + plain)
            else:
                lines.append(md)
    while lines and lines[-1] == "":
        lines.pop()
    # 연속된 빈 줄 정리
    tidy: list[str] = []
    for line in lines:
        if line == "" and tidy and tidy[-1] == "":
            continue
        tidy.append(line)
    return title, tidy


def swift_string(lines: list[str]) -> str:
    body = "\n".join(lines)
    if '"""#' in body:
        raise SystemExit("본문에 '\"\"\"#'가 있어 raw 문자열로 넣을 수 없다")
    return '#"""\n' + body + '\n"""#'


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    old = [p.text for p in docx.Document(sys.argv[1]).paragraphs]
    new_doc = docx.Document(sys.argv[2])
    new_ps = new_doc.paragraphs
    new = [p.text for p in new_ps]
    topics = question_topics()

    guides: dict[str, tuple[str, str, list[str]]] = {}
    sm = difflib.SequenceMatcher(None, old, new, autojunk=False)
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag != "insert":
            continue
        qid = None
        for k in range(i1 - 1, -1, -1):
            m = re.match(r"^(\d{4}-\d+)\b", old[k].strip())
            if m and m.group(1) in topics:
                qid = m.group(1)
                break
        if qid is None:
            raise SystemExit(f"{i1}번 문단 앞에서 문항 번호를 찾지 못했다")
        topic = topics[qid]
        title, lines = convert(new_ps[j1:j2])
        if topic in guides:
            raise SystemExit(f"주제 '{topic}'에 설명이 두 번 붙어 있다({guides[topic][0]}, {qid})")
        guides[topic] = (qid, title, lines)

    parts = []
    for topic, (qid, title, lines) in guides.items():
        parts.append(
            f"        // {qid} 아래에 붙은 설명\n"
            f"        TopicGuide(\n"
            f"            topics: {swift_array([topic] + SHARED_TOPICS.get(topic, []))},\n"
            f"            title: \"{title}\",\n"
            f"            body: {indent(swift_string(lines), 12)}\n"
            f"        ),")
    OUT_FILE.write_text(HEADER + "\n".join(parts) + FOOTER, encoding="utf-8")
    print(f"{len(guides)}개 주제 설명을 {OUT_FILE.relative_to(REPO_ROOT)}에 썼다")
    for topic, (qid, title, lines) in guides.items():
        print(f"  {qid} {topic}: {title} ({len(lines)}줄)")


def swift_array(items: list[str]) -> str:
    return "[" + ", ".join(f'"{t}"' for t in items) + "]"


def indent(s: str, n: int) -> str:
    pad = " " * n
    first, *rest = s.split("\n")
    return "\n".join([first] + [pad + line if line else "" for line in rest])


HEADER = '''import Foundation

// 이 파일은 Scripts/import_topic_guides.py가 검토본 docx에서 만들었다.
// 본문은 줄 단위의 간단한 표기를 쓴다: "# 소제목", "- 목록"(두 칸 들여쓰기마다 한 수준),
// 문장 안의 **굵게**. TopicGuideView가 이 표기를 화면에 그린다.

/// 주제별 설명. topics에 들어 있는 주제(StudyTopics의 주제 이름)에서 이 설명을 보여 준다.
struct TopicGuide: Identifiable {
    let topics: [String]
    let title: String
    let body: String
    var id: String { title }
}

enum TopicGuides {
    static func guide(for topic: String) -> TopicGuide? {
        all.first { $0.topics.contains(topic) }
    }

    static let all: [TopicGuide] = [
'''

FOOTER = '''
    ]
}
'''

if __name__ == "__main__":
    main()
